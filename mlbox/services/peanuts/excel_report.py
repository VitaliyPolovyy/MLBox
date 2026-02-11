"""
Excel report generation for peanut processing results.

Builds the .xlsx workbook (Результат, Данні, Result Image sheets) from
PeanutProcessingResult. All logic is here to keep peanuts.py focused on
processing and API.
"""
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.drawing.image import Image as OpenPyxlImage
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from mlbox.models.peanuts.cls.yolo_cls_model import YOLOPeanutsClassifier
from mlbox.services.peanuts.datatype import PeanutProcessingResult


def _fmt_num(v: float, drop_trailing_zero: bool) -> str:
    """Format number; if drop_trailing_zero, show 43 instead of 43.0."""
    if drop_trailing_zero and v == int(v):
        s = str(int(v))
    else:
        s = f"{v:.1f}"
    return s.replace(".", ",")


def _histogram_with_overflow(
    values: list,
    v_min: float,
    v_max: float,
    step: float,
    label_fmt: str,
    drop_trailing_zero: bool = False,
) -> tuple:
    """Fixed range: first bin '< min', then min..max in step, last bin '> max'."""
    edges = np.arange(v_min, v_max + step * 0.5, step)
    counts, _ = np.histogram(values, bins=edges)
    n_below = sum(1 for v in values if v < v_min)
    n_above = sum(1 for v in values if v > v_max)
    labels = [f"< {_fmt_num(v_min, drop_trailing_zero)}"]
    labels += [
        f"{_fmt_num(edges[i], drop_trailing_zero)}-{_fmt_num(edges[i + 1], drop_trailing_zero)}"
        for i in range(len(counts))
    ]
    labels.append(f"> {_fmt_num(v_max, drop_trailing_zero)}")
    counts_with_overflow = [n_below] + list(counts) + [n_above]
    return labels, counts_with_overflow


def prepare_excel(
    peanut_processing_result: PeanutProcessingResult,
    output_dir: Path,
) -> Path:
    """
    Create an Excel file for the given PeanutProcessingResult.

    Args:
        peanut_processing_result: The result to prepare the Excel file for.
        output_dir: Directory where the .xlsx file will be saved.

    Returns:
        Path to the generated Excel file.
    """
    excel_file = output_dir / f"{Path(peanut_processing_result.original_image_filename).stem}.xlsx"

    if excel_file.exists():
        excel_file.unlink()

    with pd.ExcelWriter(excel_file, engine="openpyxl") as writer:
        writer.book.create_sheet(title="Результат", index=0)
        writer.book.active = 0

        result_table = pd.DataFrame(
            [
                {
                    "№ п/п": idx,
                    "max діаметр, мм": (
                        round(
                            peanut.ellipse.axes[1] / peanut_processing_result.pixels_per_mm,
                            1,
                        )
                        if peanut.ellipse
                        else None
                    ),
                    "min діаметр, мм": (
                        round(
                            peanut.ellipse.axes[0] / peanut_processing_result.pixels_per_mm,
                            1,
                        )
                        if peanut.ellipse
                        else None
                    ),
                    "площа маски, мм²": (
                        round(
                            np.sum(peanut.mask) / (peanut_processing_result.pixels_per_mm**2),
                            2,
                        )
                        if peanut.mask is not None
                        else None
                    ),
                    "клас": YOLOPeanutsClassifier.class_names[peanut.real_class[0]],
                    "впевненність (клас)": round(peanut.real_class[1], 2),
                    "впевненність (маска)": round(peanut.det_confidence, 2),
                }
                for idx, peanut in enumerate(peanut_processing_result.peanuts)
            ]
        )

        major_axes_mm = [
            peanut.ellipse.axes[1] / peanut_processing_result.pixels_per_mm
            for peanut in peanut_processing_result.peanuts
            if peanut.ellipse
        ]
        minor_axes_mm = [
            peanut.ellipse.axes[0] / peanut_processing_result.pixels_per_mm
            for peanut in peanut_processing_result.peanuts
            if peanut.ellipse
        ]
        areas_mm2 = [
            np.sum(peanut.mask) / (peanut_processing_result.pixels_per_mm**2)
            for peanut in peanut_processing_result.peanuts
            if peanut.mask is not None
        ]

        histogram_minor_df = None
        histogram_major_df = None
        histogram_area_df = None
        minor_chart = None
        major_chart = None
        area_chart = None

        if peanut_processing_result.peanuts and any(
            peanut.ellipse for peanut in peanut_processing_result.peanuts
        ):
            AREA_MIN, AREA_MAX, AREA_STEP = 43.0, 133.0, 5.0
            MAJOR_MIN, MAJOR_MAX, MAJOR_STEP = 8.5, 17.0, 0.5
            MINOR_MIN, MINOR_MAX, MINOR_STEP = 6.0, 11.0, 0.5

            if minor_axes_mm:
                labels_minor, counts_minor = _histogram_with_overflow(
                    minor_axes_mm, MINOR_MIN, MINOR_MAX, MINOR_STEP, ""
                )
                histogram_minor_data = [
                    {"мітка (меньша вісь)": lb, "кількість (меньша вісь)": int(c)}
                    for lb, c in zip(labels_minor, counts_minor)
                ]
                histogram_minor_df = pd.DataFrame(histogram_minor_data)
                minor_chart = BarChart()
                minor_chart.type = "col"
                minor_chart.style = 10
                minor_chart.title = "Розподіл частот меншої осі арахісу"
                minor_chart.y_axis.title = "Частота, шт"
                minor_chart.x_axis.delete = False
                minor_chart.legend = None
                minor_chart.width = 12.75
                minor_chart.height = 8.5

            if major_axes_mm:
                labels_major, counts_major = _histogram_with_overflow(
                    major_axes_mm, MAJOR_MIN, MAJOR_MAX, MAJOR_STEP, ""
                )
                histogram_major_data = [
                    {"мітка (більша вісь)": lb, "кількість (більша вісь)": int(c)}
                    for lb, c in zip(labels_major, counts_major)
                ]
                histogram_major_df = pd.DataFrame(histogram_major_data)
                major_chart = BarChart()
                major_chart.type = "col"
                major_chart.style = 10
                major_chart.title = "Розподіл частот більшої осі арахісу"
                major_chart.y_axis.title = "Частота, шт"
                major_chart.x_axis.delete = False
                major_chart.legend = None
                major_chart.width = 12.75
                major_chart.height = 8.5

            if areas_mm2:
                labels_area, counts_area = _histogram_with_overflow(
                    areas_mm2, AREA_MIN, AREA_MAX, AREA_STEP, "", drop_trailing_zero=True
                )
                histogram_area_data = [
                    {"мітка (площа)": lb, "кількість (площа)": int(c)}
                    for lb, c in zip(labels_area, counts_area)
                ]
                histogram_area_df = pd.DataFrame(histogram_area_data)
                area_chart = BarChart()
                area_chart.type = "col"
                area_chart.style = 10
                area_chart.title = "Розподіл частот площі арахісу"
                area_chart.y_axis.title = "Частота, шт"
                area_chart.x_axis.delete = False
                area_chart.legend = None
                area_chart.width = 12.75
                area_chart.height = 8.5

        result_table.to_excel(writer, sheet_name="Данні", index=False, startrow=0)
        data_sheet = writer.sheets["Данні"]
        n_peanuts = len(peanut_processing_result.peanuts)

        indicators_table_start_col = result_table.shape[1] + 2
        indicators_table_num_cols = 4
        ind_start_row = 2
        row_names_uk = [
            "Мін",
            "Макс",
            "Розмах",
            "Мода",
            "Медіана",
            "Ср знач ",
            "Ст кв відхилення",
            "Коефіцієнт варіації",
        ]
        major_s = result_table["max діаметр, мм"].dropna()
        minor_s = result_table["min діаметр, мм"].dropna()
        area_s = result_table["площа маски, мм²"].dropna()

        def _stats(series: pd.Series, decimals: int = 2) -> list:
            if series.empty:
                return [None] * 8
            s = series.astype(float)
            mn, mx = s.min(), s.max()
            mode_val = s.mode().iloc[0] if len(s.mode()) else None
            med = s.median()
            mean = s.mean()
            stdev = s.std(ddof=1) if len(s) >= 2 else (0.0 if len(s) == 1 else None)
            cv = (
                (stdev / mean * 100)
                if (mean and mean != 0 and stdev is not None)
                else None
            )
            round_ = (
                lambda x: round(x, decimals)
                if x is not None and isinstance(x, (int, float))
                else x
            )
            return [
                round_(mn),
                round_(mx),
                round_(mx - mn),
                round_(mode_val),
                round_(med),
                round_(mean),
                round_(stdev),
                round_(cv),
            ]

        ind_major = _stats(major_s, decimals=1)
        ind_minor = _stats(minor_s, decimals=1)
        ind_area = _stats(area_s, decimals=2)

        data_sheet.cell(row=1, column=indicators_table_start_col, value="Показник").font = Font(
            bold=True
        )
        data_sheet.cell(
            row=1, column=indicators_table_start_col + 1, value="Більша вісь"
        ).font = Font(bold=True)
        data_sheet.cell(
            row=1, column=indicators_table_start_col + 2, value="Менша вісь"
        ).font = Font(bold=True)
        data_sheet.cell(
            row=1, column=indicators_table_start_col + 3, value="Площа арахісу"
        ).font = Font(bold=True)
        for i, name in enumerate(row_names_uk):
            r = ind_start_row + i
            data_sheet.cell(row=r, column=indicators_table_start_col, value=name)
            data_sheet.cell(row=r, column=indicators_table_start_col + 1, value=ind_major[i])
            data_sheet.cell(row=r, column=indicators_table_start_col + 2, value=ind_minor[i])
            data_sheet.cell(row=r, column=indicators_table_start_col + 3, value=ind_area[i])
        hist_start_row = 1
        minor_start_col = 14
        major_start_col = 17
        area_start_col = 20

        if histogram_minor_df is not None:
            for c_idx, col_name in enumerate(histogram_minor_df.columns, start=1):
                cell = data_sheet.cell(
                    row=hist_start_row, column=minor_start_col + c_idx - 1, value=col_name
                )
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.font = Font(bold=True)
            for r_idx, row in enumerate(histogram_minor_df.values, start=1):
                for c_idx, value in enumerate(row, start=1):
                    cell = data_sheet.cell(
                        row=hist_start_row + r_idx,
                        column=minor_start_col + c_idx - 1,
                        value=value,
                    )
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            if minor_chart is not None:
                minor_chart.add_data(
                    Reference(
                        data_sheet,
                        min_col=minor_start_col + 1,
                        min_row=hist_start_row,
                        max_row=hist_start_row + len(histogram_minor_df),
                    ),
                    titles_from_data=True,
                )
                minor_chart.set_categories(
                    Reference(
                        data_sheet,
                        min_col=minor_start_col,
                        min_row=hist_start_row + 1,
                        max_row=hist_start_row + len(histogram_minor_df),
                    )
                )
                if minor_chart.series:
                    s1 = minor_chart.series[0]
                    s1.graphicalProperties.solidFill = "4472C4"
                    s1.graphicalProperties.line.solidFill = "4472C4"
                minor_chart.dataLabels = DataLabelList()
                minor_chart.dataLabels.showVal = True
                minor_chart.dataLabels.showCatName = False
                minor_chart.dataLabels.showSerName = False
                minor_chart.dataLabels.showLegendKey = False
                minor_chart.dataLabels.showLeaderLines = False
                minor_chart.dataLabels.position = "outEnd"
        if histogram_major_df is not None:
            for c_idx, col_name in enumerate(histogram_major_df.columns, start=1):
                cell = data_sheet.cell(
                    row=hist_start_row, column=major_start_col + c_idx - 1, value=col_name
                )
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.font = Font(bold=True)
            for r_idx, row in enumerate(histogram_major_df.values, start=1):
                for c_idx, value in enumerate(row, start=1):
                    cell = data_sheet.cell(
                        row=hist_start_row + r_idx,
                        column=major_start_col + c_idx - 1,
                        value=value,
                    )
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            if major_chart is not None:
                major_chart.add_data(
                    Reference(
                        data_sheet,
                        min_col=major_start_col + 1,
                        min_row=hist_start_row,
                        max_row=hist_start_row + len(histogram_major_df),
                    ),
                    titles_from_data=True,
                )
                major_chart.set_categories(
                    Reference(
                        data_sheet,
                        min_col=major_start_col,
                        min_row=hist_start_row + 1,
                        max_row=hist_start_row + len(histogram_major_df),
                    )
                )
                if major_chart.series:
                    s1 = major_chart.series[0]
                    s1.graphicalProperties.solidFill = "4472C4"
                    s1.graphicalProperties.line.solidFill = "4472C4"
                major_chart.dataLabels = DataLabelList()
                major_chart.dataLabels.showVal = True
                major_chart.dataLabels.showCatName = False
                major_chart.dataLabels.showSerName = False
                major_chart.dataLabels.showLegendKey = False
                major_chart.dataLabels.showLeaderLines = False
                major_chart.dataLabels.position = "outEnd"
        if histogram_area_df is not None:
            for c_idx, col_name in enumerate(histogram_area_df.columns, start=1):
                cell = data_sheet.cell(
                    row=hist_start_row, column=area_start_col + c_idx - 1, value=col_name
                )
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.font = Font(bold=True)
            for r_idx, row in enumerate(histogram_area_df.values, start=1):
                for c_idx, value in enumerate(row, start=1):
                    cell = data_sheet.cell(
                        row=hist_start_row + r_idx,
                        column=area_start_col + c_idx - 1,
                        value=value,
                    )
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            if area_chart is not None:
                area_chart.add_data(
                    Reference(
                        data_sheet,
                        min_col=area_start_col + 1,
                        min_row=hist_start_row,
                        max_row=hist_start_row + len(histogram_area_df),
                    ),
                    titles_from_data=True,
                )
                area_chart.set_categories(
                    Reference(
                        data_sheet,
                        min_col=area_start_col,
                        min_row=hist_start_row + 1,
                        max_row=hist_start_row + len(histogram_area_df),
                    )
                )
                if area_chart.series:
                    s1 = area_chart.series[0]
                    s1.graphicalProperties.solidFill = "4472C4"
                    s1.graphicalProperties.line.solidFill = "4472C4"
                area_chart.dataLabels = DataLabelList()
                area_chart.dataLabels.showVal = True
                area_chart.dataLabels.showCatName = False
                area_chart.dataLabels.showSerName = False
                area_chart.dataLabels.showLegendKey = False
                area_chart.dataLabels.showLeaderLines = False
                area_chart.dataLabels.position = "outEnd"

        result_sheet = writer.sheets["Результат"]
        chart_row = 1
        chart_cols_per_histogram = 8
        current_col_num = 1
        if minor_chart is not None:
            result_sheet.add_chart(
                minor_chart, f"{get_column_letter(current_col_num)}{chart_row}"
            )
            current_col_num += chart_cols_per_histogram
        if major_chart is not None:
            result_sheet.add_chart(
                major_chart, f"{get_column_letter(current_col_num)}{chart_row}"
            )
            current_col_num += chart_cols_per_histogram
        if area_chart is not None:
            result_sheet.add_chart(
                area_chart, f"{get_column_letter(current_col_num)}{chart_row}"
            )

        summary_table_start_row = 22
        data_sheet_name = "Данні"
        cv_menша_вісь_cell = f"'{data_sheet_name}'!K9"
        pieces_per_ounce = round(len(peanut_processing_result.peanuts) * 28.35 / 100, 2)
        summary_indicators = [
            (
                "Однорідність",
                f'=IF({cv_menша_вісь_cell}<=18,"Однорідний","Неоднорідний")',
            ),
            ("Шт в 1 унції", pieces_per_ounce),
            ("Вага зразку, г", peanut_processing_result.weight_g),
            ("Вибірка", peanut_processing_result.sample_size),
        ]
        result_sheet.cell(
            row=summary_table_start_row, column=1, value="Показник"
        ).font = Font(bold=True)
        result_sheet.cell(
            row=summary_table_start_row, column=2, value="Значення"
        ).font = Font(bold=True)
        for i, (name, val) in enumerate(summary_indicators, start=1):
            result_sheet.cell(row=summary_table_start_row + i, column=1, value=name)
            cell_b2 = result_sheet.cell(row=summary_table_start_row + i, column=2)
            if isinstance(val, str) and val.startswith("="):
                cell_b2.value = val
            else:
                cell_b2.value = val
        for r in range(
            summary_table_start_row,
            summary_table_start_row + len(summary_indicators) + 1,
        ):
            for c in (1, 2):
                result_sheet.cell(row=r, column=c).alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                    horizontal="left" if c == 1 else "center",
                )
        result_sheet.column_dimensions["A"].width = 22
        result_sheet.column_dimensions["B"].width = 14

        worksheet = writer.sheets["Данні"]
        worksheet.column_dimensions[
            get_column_letter(indicators_table_start_col)
        ].width = 22
        for c in range(
            indicators_table_start_col + 1,
            indicators_table_start_col + indicators_table_num_cols,
        ):
            worksheet.column_dimensions[get_column_letter(c)].width = 14
        histogram_label_width = 22
        worksheet.column_dimensions[get_column_letter(minor_start_col)].width = (
            histogram_label_width
        )
        worksheet.column_dimensions[get_column_letter(major_start_col)].width = (
            histogram_label_width
        )
        worksheet.column_dimensions[get_column_letter(area_start_col)].width = (
            histogram_label_width * 2
        )

        caption_alignment = Alignment(
            wrap_text=True, horizontal="center", vertical="center"
        )
        caption_columns = (
            list(range(1, 8))
            + list(
                range(
                    indicators_table_start_col,
                    indicators_table_start_col + indicators_table_num_cols,
                )
            )
            + list(range(minor_start_col, minor_start_col + 2))
            + list(range(major_start_col, major_start_col + 2))
            + list(range(area_start_col, area_start_col + 2))
        )

        for row in worksheet.iter_rows():
            for cell in row:
                if cell.row == 1 and cell.column in caption_columns:
                    cell.alignment = caption_alignment
                elif cell.column == indicators_table_start_col:
                    cell.alignment = Alignment(
                        wrap_text=True, horizontal="left", vertical="center"
                    )
                elif (
                    indicators_table_start_col
                    < cell.column
                    <= indicators_table_start_col + indicators_table_num_cols - 1
                ):
                    cell.alignment = Alignment(
                        wrap_text=True, horizontal="center", vertical="center"
                    )
                else:
                    cell.alignment = Alignment(
                        wrap_text=True, horizontal="center", vertical="center"
                    )

        result_image_sheet = writer.book.create_sheet(title="Result Image")
        image_stream = BytesIO()
        peanut_processing_result.result_image.save(image_stream, format="PNG")
        image_stream.seek(0)
        openpyxl_image = OpenPyxlImage(image_stream)
        openpyxl_image.width = openpyxl_image.width // 2
        openpyxl_image.height = openpyxl_image.height // 2
        result_image_sheet.add_image(openpyxl_image, "A1")

    return excel_file
